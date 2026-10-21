"""Таблицы из Excel/CSV, присланных ассистенту в личку.

21.10.2026: Total.xlsx владельца (9 листов, в «Total» 11 001 строка) не помещался
в текстовый лимит вложения — бот видел только первые листы и просил резать файл.
Читать тысячи строк текстом модель всё равно не может. Поэтому каждый лист
становится таблицей в отдельной маленькой SQLite-базе (одна на вложение), модель
получает схему (листы, колонки с буквами, первые строки) и считает SQL-запросами —
инструмент query_attachment. Запросы только на чтение, к одной этой базе.
"""
from __future__ import annotations

import csv
import io
import json
import os
import re
import sqlite3
import time
from datetime import date, datetime

import database as db

MAX_ROWS_PER_SHEET = 200_000
MAX_COLS = 200
HEADER_SCAN = 10
QUERY_TIMEOUT_S = 8
QUERY_MAX_ROWS = 1000
PREVIEW_ROWS = 3
TABLE_EXTS = {"xlsx", "xlsm", "xls", "csv", "tsv"}


def tables_dir() -> str:
    explicit = (os.getenv("AI_ATTACHMENT_TABLES_DIR") or "").strip()
    if explicit:
        return explicit
    return os.path.join(os.path.dirname(os.path.abspath(db.DB_PATH)) or ".", "ai_attachments")


def path_for(att_id) -> str:
    return os.path.join(tables_dir(), f"{int(att_id)}.sqlite")


def has_tables(att_id) -> bool:
    try:
        return os.path.exists(path_for(att_id))
    except (TypeError, ValueError):
        return False


def _letter(i: int) -> str:
    s = ""
    i += 1
    while i:
        i, r = divmod(i - 1, 26)
        s = chr(65 + r) + s
    return s


def _cell(v):
    if v is None:
        return None
    if isinstance(v, bool):
        return int(v)
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d %H:%M") if (v.hour or v.minute) else v.strftime("%Y-%m-%d")
    if isinstance(v, date):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, (int, float)):
        return v
    s = str(v).strip()
    return s if s else None


def _pick_header(rows: list) -> int | None:
    """Строка заголовка: среди первых строк — та, где больше всего непустых текстовых
    ячеек (у «Ombor» над шапкой есть строка групп «MIJOZ / YUK NOMI…»)."""
    best, best_n = None, 1
    for idx, row in enumerate(rows[:HEADER_SCAN]):
        n = sum(1 for v in row if isinstance(v, str) and v.strip() and not re.fullmatch(r"[-\d.,\s:]+", v))
        if n > best_n:
            best, best_n = idx, n
    return best


def _column_names(header: list, width: int) -> list:
    names, seen = [], set()
    for i in range(width):
        raw = header[i] if header and i < len(header) else None
        name = re.sub(r"\s+", " ", str(raw)).strip() if raw not in (None, "") else ""
        name = name[:60] or _letter(i)
        base, n = name, 2
        while name.lower() in seen or name.lower() == "_row":
            name = f"{base}_{n}"
            n += 1
        seen.add(name.lower())
        names.append({"name": name, "letter": _letter(i)})
    return names


def _sheet(name: str, rows: list, first_row_number: int = 1) -> dict | None:
    rows = [[_cell(v) for v in r] for r in rows]
    width = 0
    for r in rows:
        for i in range(len(r) - 1, -1, -1):
            if r[i] is not None:
                width = max(width, i + 1)
                break
    width = min(width, MAX_COLS)
    if not width:
        return None
    h = _pick_header(rows)
    header = rows[h] if h is not None else None
    start = (h + 1) if h is not None else 0
    data, numbers, truncated = [], [], False
    for idx in range(start, len(rows)):
        r = rows[idx][:width]
        if not any(v is not None for v in r):
            continue
        if len(data) >= MAX_ROWS_PER_SHEET:
            truncated = True
            break
        data.append(r + [None] * (width - len(r)))
        numbers.append(first_row_number + idx)
    return {"sheet": name, "header_row": (first_row_number + h) if h is not None else 0,
            "columns": _column_names(header, width), "rows": data, "row_numbers": numbers, "truncated": truncated}


def extract_tables(data: bytes, filename: str) -> list:
    """Листы файла → [{sheet, header_row, columns[{name, letter}], rows, row_numbers, truncated}]."""
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    out = []
    if ext in ("xlsx", "xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            for ws in wb.worksheets:
                rows = []
                for r in ws.iter_rows(values_only=True):
                    rows.append(list(r))
                    if len(rows) > MAX_ROWS_PER_SHEET + HEADER_SCAN:
                        break
                t = _sheet(ws.title, rows)
                if t:
                    out.append(t)
        finally:
            wb.close()
    elif ext == "xls":
        import xlrd
        book = xlrd.open_workbook(file_contents=data)
        for sh in book.sheets():
            rows = []
            for r in range(min(sh.nrows, MAX_ROWS_PER_SHEET + HEADER_SCAN)):
                row = []
                for c in range(sh.ncols):
                    cell = sh.cell(r, c)
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        try:
                            row.append(xlrd.xldate.xldate_as_datetime(cell.value, book.datemode))
                            continue
                        except Exception:
                            pass
                    row.append(cell.value if cell.value != "" else None)
                rows.append(row)
            t = _sheet(sh.name, rows)
            if t:
                out.append(t)
    elif ext in ("csv", "tsv"):
        from services.attachment_reader import decode_text
        text = decode_text(data)
        try:
            dialect = csv.Sniffer().sniff(text[:5000], delimiters=",;\t|")
            delim = dialect.delimiter
        except csv.Error:
            delim = "\t" if ext == "tsv" else ","
        rows = []
        for r in csv.reader(io.StringIO(text), delimiter=delim):
            rows.append([_num(v) for v in r])
            if len(rows) > MAX_ROWS_PER_SHEET + HEADER_SCAN:
                break
        t = _sheet(filename.rsplit(".", 1)[0][:40] or "csv", rows)
        if t:
            out.append(t)
    return out


def _num(v: str):
    s = str(v).strip()
    if re.fullmatch(r"-?\d+", s):
        try:
            return int(s)
        except ValueError:
            return s
    if re.fullmatch(r"-?\d+[.,]\d+", s):
        try:
            return float(s.replace(",", "."))
        except ValueError:
            return s
    return s or None


def _q(name: str) -> str:
    return '"' + str(name).replace('"', '""') + '"'


def _table_name(sheet: str, used: set) -> str:
    base = re.sub(r"\s+", " ", str(sheet or "")).strip()[:60] or "sheet"
    name, n = base, 2
    while name.lower() in used or name.lower().startswith("sqlite_") or name.lower() == "__sheets":
        name = f"{base}_{n}"
        n += 1
    used.add(name.lower())
    return name


def store(att_id, tables: list) -> list:
    """Записать листы в базу вложения; вернуть краткую схему [{sheet, table, rows, …}]."""
    os.makedirs(tables_dir(), exist_ok=True)
    path = path_for(att_id)
    tmp = path + ".tmp"
    if os.path.exists(tmp):
        os.remove(tmp)
    conn = sqlite3.connect(tmp)
    summary, used = [], set()
    try:
        conn.execute("CREATE TABLE __sheets (sheet TEXT, table_name TEXT, header_row INTEGER, rows INTEGER, "
                     "columns TEXT, truncated INTEGER)")
        for t in tables:
            tname = _table_name(t["sheet"], used)
            cols = ["_row"] + [c["name"] for c in t["columns"]]
            conn.execute(f"CREATE TABLE {_q(tname)} ({', '.join(_q(c) for c in cols)})")
            marks = ", ".join("?" for _ in cols)
            conn.executemany(f"INSERT INTO {_q(tname)} VALUES ({marks})",
                             ([n] + r for n, r in zip(t["row_numbers"], t["rows"])))
            conn.execute("INSERT INTO __sheets VALUES (?, ?, ?, ?, ?, ?)",
                         (t["sheet"], tname, t["header_row"], len(t["rows"]),
                          json.dumps(t["columns"], ensure_ascii=False), int(t["truncated"])))
            summary.append({"sheet": t["sheet"], "table": tname, "rows": len(t["rows"]),
                            "header_row": t["header_row"], "columns": t["columns"], "truncated": t["truncated"],
                            "sample": t["rows"][:PREVIEW_ROWS]})
        conn.commit()
    finally:
        conn.close()
    os.replace(tmp, path)
    return summary


def _short(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float):
        v = f"{v:.4f}".rstrip("0").rstrip(".")
    s = str(v).replace("\n", " ")
    return s if len(s) <= 40 else s[:40] + "…"


def schema_text(tables: list) -> str:
    """Схема для модели: листы → таблицы, колонки с буквами, первые строки."""
    lines = []
    for t in tables:
        rows = t.get("rows") if "rows" in t and isinstance(t.get("rows"), list) else None
        count = len(rows) if rows is not None else t.get("rows")
        sample = (rows or [])[:PREVIEW_ROWS] if rows is not None else t.get("sample") or []
        tname = t.get("table") or t["sheet"]
        head = (f"## Лист «{t['sheet']}» → таблица {_q(tname)}: {count} строк данных"
                + (f", заголовок в строке {t['header_row']}" if t.get("header_row") else ", без заголовка")
                + (f" (обрезано до {MAX_ROWS_PER_SHEET})" if t.get("truncated") else ""))
        cols = ", ".join(f"{c['letter']} {_q(c['name'])}" for c in t["columns"])
        lines.append(head)
        lines.append(f"Колонки: _row (номер строки в Excel), {cols}")
        for r in sample:
            lines.append("  " + " | ".join(_short(v) for v in r))
    return "\n".join(lines)


_ALLOWED_ACTIONS = {sqlite3.SQLITE_SELECT, sqlite3.SQLITE_READ, sqlite3.SQLITE_FUNCTION,
                    getattr(sqlite3, "SQLITE_RECURSIVE", 33)}


def _authorizer(action, arg1, arg2, dbname, source):
    if action == sqlite3.SQLITE_FUNCTION and str(arg2 or "").lower() in ("load_extension", "readfile", "writefile"):
        return sqlite3.SQLITE_DENY
    return sqlite3.SQLITE_OK if action in _ALLOWED_ACTIONS else sqlite3.SQLITE_DENY


def query(att_id, sql: str, limit: int = 200) -> dict:
    """Один SELECT к базе вложения: только чтение, без ATTACH/PRAGMA, с таймаутом."""
    sql = str(sql or "").strip().rstrip(";").strip()
    if not re.match(r"(?is)^(select|with)\b", sql):
        return {"error": "Нужен один запрос SELECT (или WITH … SELECT)"}
    if ";" in sql:
        return {"error": "Только один запрос за раз"}
    path = path_for(att_id)
    if not os.path.exists(path):
        return {"error": "Таблиц у этого вложения нет — это не Excel/CSV или файл старше хранения"}
    limit = max(1, min(int(limit or 200), QUERY_MAX_ROWS))
    conn = sqlite3.connect(f"file:{path}?mode=ro", uri=True)
    deadline = time.time() + QUERY_TIMEOUT_S
    conn.set_progress_handler(lambda: 1 if time.time() > deadline else 0, 20000)
    conn.set_authorizer(_authorizer)
    try:
        cur = conn.execute(sql)
        cols = [d[0] for d in (cur.description or [])]
        rows = cur.fetchmany(limit + 1)
    except sqlite3.DatabaseError as exc:
        msg = str(exc)
        if "interrupted" in msg.lower():
            msg = f"запрос дольше {QUERY_TIMEOUT_S} с — сузь его (WHERE, GROUP BY)"
        return {"error": f"SQL: {msg}"}
    finally:
        conn.close()
    return {"columns": cols, "rows": [list(r) for r in rows[:limit]], "returned": min(len(rows), limit),
            "truncated": len(rows) > limit}


def prune(keep_ids) -> None:
    """Удалить базы вложений, которых больше нет в ai_attachments (хранятся последние 10 на человека)."""
    keep = {str(int(i)) for i in keep_ids or []}
    folder = tables_dir()
    if not os.path.isdir(folder):
        return
    for name in os.listdir(folder):
        m = re.fullmatch(r"(\d+)\.sqlite(\.tmp)?", name)
        if m and m.group(1) not in keep:
            try:
                os.remove(os.path.join(folder, name))
            except OSError:
                pass
